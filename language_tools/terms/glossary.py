"""Read/write a glossary file (csv/xlsx) of ``TermEntry`` rows
(DESIGN.md section 15.1, Phase G0).

Column shape is deliberately different from the bilingual-source readers
in ``language_tools/readers/``: those *guess* a header/column layout
because their input is an arbitrary externally-authored document. A
glossary file here is either round-tripped by ``write()`` in this module
or hand-edited starting from that output, so requiring a *named* header
(``src_term``/``tgt_term``/...) rather than guessing column positions is
the right tradeoff -- position-guessing a term list would also misfire on
``readers/_rowreader.py``'s own ``looks_like_header()`` heuristic (short
cells, no sentence punctuation): *every* row in a term list looks like
that, header or not, so that heuristic can't tell header from data here.

Language pair is a file-level property, not a per-row column: a glossary
in this tool covers one src/tgt language pair (same mental model as a
TM), so asking someone hand-maintaining the file to repeat "en-US,zh-CN"
on every single row would be pure busywork the file format doesn't need.
``read()`` takes the pair as parameters and stamps every entry with it;
``write()`` doesn't persist it at all (entries already carry their own
``src_lang``/``tgt_lang`` -- see model.py -- ``write()`` just doesn't
duplicate that into a column, since a single glossary file is always one
pair by convention here).

Encoding fallback for csv mirrors ``readers/csv_bilingual.py`` exactly
(utf-8-sig -> utf-8 -> gb18030) -- same real-world "Excel export from
Chinese Windows isn't UTF-8" problem, same fix, not reinvented.

``import openpyxl`` is deferred to inside the two functions that actually
need it, same reasoning and same fix as ``readers/xlsx_bilingual.py``'s
docstring: this module (like every ``toolbox/tools/*/page.py``) gets
imported at app startup regardless of whether the person ever opens an
xlsx glossary in the session, so an eager import would make every launch
pay openpyxl's load cost, not just the ones that use it.
"""
import csv
import io
import os

from language_tools.terms.model import TermEntry

ENCODINGS = ['utf-8-sig', 'utf-8', 'gb18030']

# Header column names, in write order. Read matches header cells
# case-insensitively/whitespace-trimmed against these; write always
# emits exactly this set, in this order.
COLUMNS = ['src_term', 'tgt_term', 'status', 'domain', 'note']

_VALID_STATUSES = {'approved', 'forbidden'}
_SUPPORTED_EXTS = ('.csv', '.xlsx', '.xlsm')


def _read_text(path):
    last_err = None
    for enc in ENCODINGS:
        try:
            with open(path, encoding=enc) as f:
                return f.read(), enc
        except UnicodeDecodeError as e:
            last_err = e
    raise ValueError('could not decode %s with any of %s: %s' % (path, ENCODINGS, last_err))


def _parse_header(row):
    """Maps recognized column names to their position. Requires
    src_term/tgt_term at minimum -- a glossary without those two has
    nothing to check against, so failing loudly here (rather than
    returning an empty entry list a caller might mistake for "empty
    file") is the right failure mode.
    """
    col_index = {}
    for i, cell in enumerate(row):
        name = (cell or '').strip().lower()
        if name in COLUMNS:
            col_index[name] = i
    missing = [c for c in ('src_term', 'tgt_term') if c not in col_index]
    if missing:
        raise ValueError(
            '术语表缺少必需的表头列 %s（需要 src_term/tgt_term 表头，大小写不敏感）' % missing)
    return col_index


def _cell(cells, col_index, name):
    idx = col_index.get(name)
    if idx is None or idx >= len(cells):
        return ''
    return (cells[idx] or '').strip()


def _rows_to_entries(rows, src_lang, tgt_lang):
    """``rows`` is a list of raw cell-lists, row 0 assumed to be the
    header. Returns [] for a header-only or empty file, same "no data,
    not an error" treatment as the bilingual readers give an empty
    source document.
    """
    if not rows:
        return []
    col_index = _parse_header(rows[0])

    entries, bad_status_rows = [], []
    for row_number, cells in enumerate(rows[1:], 2):
        src_term = _cell(cells, col_index, 'src_term')
        tgt_term = _cell(cells, col_index, 'tgt_term')
        if not src_term and not tgt_term:
            continue
        status = _cell(cells, col_index, 'status').lower() or 'approved'
        if status not in _VALID_STATUSES:
            bad_status_rows.append(row_number)
            status = 'approved'
        entries.append(TermEntry(
            src_lang=src_lang, tgt_lang=tgt_lang, src_term=src_term, tgt_term=tgt_term,
            status=status, domain=_cell(cells, col_index, 'domain') or None,
            note=_cell(cells, col_index, 'note') or None,
        ))
    if bad_status_rows:
        print('warning: glossary rows with unrecognized status (treated as \'approved\'): %s'
              % bad_status_rows)
    return entries


def read(path, src_lang, tgt_lang):
    ext = os.path.splitext(path)[1].lower()
    if ext == '.csv':
        text, _enc = _read_text(path)
        rows = [row for row in csv.reader(io.StringIO(text)) if any(c.strip() for c in row)]
    elif ext in ('.xlsx', '.xlsm'):
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        try:
            ws = wb.worksheets[0]
            rows = [['' if c is None else str(c).strip() for c in row]
                    for row in ws.iter_rows(values_only=True)]
            rows = [row for row in rows if any(c for c in row)]
        finally:
            wb.close()
    else:
        raise ValueError('unsupported glossary format %r (expected one of %s)'
                          % (ext, _SUPPORTED_EXTS))
    return _rows_to_entries(rows, src_lang, tgt_lang)


def write(path, entries):
    ext = os.path.splitext(path)[1].lower()
    rows = [COLUMNS] + [
        [e.src_term, e.tgt_term, e.status, e.domain or '', e.note or ''] for e in entries
    ]
    if ext == '.csv':
        with open(path, 'w', encoding='utf-8-sig', newline='') as f:
            csv.writer(f).writerows(rows)
    elif ext in ('.xlsx', '.xlsm'):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        for row in rows:
            ws.append(row)
        wb.save(path)
    else:
        raise ValueError('unsupported glossary format %r (expected one of %s)'
                          % (ext, _SUPPORTED_EXTS))
